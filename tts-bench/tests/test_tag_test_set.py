from __future__ import annotations

import csv
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from runner.style_map import StyleMapError
from runner.tag_test_set import NEUTRAL, build_rows, detect_emotion, load_provider_tags, main

SHEET = [
    ["Category", "Common Tag", "Chirp 3 HD", "ElevenLabs v3", "Gemini 3.8 Flash TTS", "Cartesia Sonic-3",
     "Sarvam (Bulbul)", "Soniox TTS v2"],
    ["Emotion", "sad", "≈ slower", "[sad]", 'style: "sad"', '<emotion value="sad"/>', "≈ pace", "[sad]"],
    ["Emotion", "curious", "≈ question (?)", "[curious]", 'style: "curious"', "—", "≈ wording", "[curious]"],
    ["Emotion", "warm", "—", "[warmly]", 'style: "warm"', '<emotion value="affectionate"/>', "≈", "[warm]"],
]  # fmt: skip


@pytest.fixture
def tags_xlsx(tmp_path: Path) -> Path:
    wb = Workbook()
    for r in SHEET:
        wb.active.append(r)
    path = tmp_path / "tags.xlsx"
    wb.save(path)
    return path


def _csv(path: Path, rows: list[dict]) -> Path:
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    return path


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("এত কষ্ট আর সহ্য হয় না।", "sad"),
        ("তুমি কখন আসবে?", "curious"),
        ("ধন্যবাদ।", "warm"),
        ("কী সুন্দর দৃশ্য!", "surprised"),
        ("আমি ভেবেছিলাম... না, থাক।", "hesitant"),
        ("উফ্‌, কী গরম!", "annoyed"),
        ("আমি কিছু বলিনি।", NEUTRAL),  # "কি" inside "কিছু" is not a question word
        ("উৎসব আর উত্‌সব একই শব্দের দুই বানান।", NEUTRAL),
        ("হ্যাঁ।", NEUTRAL),
    ],
)
def test_detect_emotion(text: str, expected: str) -> None:
    assert detect_emotion(text) == expected


def test_provider_tags_skip_approx_dash_and_disabled_providers(tags_xlsx: Path) -> None:
    tags = load_provider_tags(tags_xlsx)
    assert tags["sad"] == {
        "elevenlabs_v3": "[sad]",
        "gemini_3.8_flash_tts": 'style: "sad"',
        "cartesia_sonic-3": '<emotion value="sad"/>',
        "soniox_tts_v2": "[sad]",
    }
    assert "cartesia_sonic-3" not in tags["curious"]  # "—" in the sheet


def test_build_rows_tags_text_and_keeps_it_exact(tmp_path: Path, tags_xlsx: Path) -> None:
    raw = " এত কষ্ট আর সহ্য হয় না। "  # surrounding spaces must survive
    data = _csv(
        tmp_path / "d.csv",
        [
            {"id": "A", "text": raw},
            {"id": "B", "text": "হ্যাঁ।"},
            {"id": "C", "text": "কেমন আছেন?"},
            {"id": "D", "text": "   "},  # blank rows are skipped
        ],
    )
    header, rows = build_rows(data, load_provider_tags(tags_xlsx))
    assert header == ["common tag", "Eleven Labs", "Gemini 3.8 Flash", "Cartesia Sonic 3", "Soniox"]
    assert rows[0] == [
        "sad",
        f"[sad] {raw}",
        f'style: "sad" {raw}',
        f'<emotion value="sad"/> {raw}',
        f"[sad] {raw}",
    ]
    assert rows[1] == [NEUTRAL, "হ্যাঁ।", "হ্যাঁ।", "হ্যাঁ।", "হ্যাঁ।"]
    assert rows[2][3] == "কেমন আছেন?"  # provider without a tag gets plain text
    assert len(rows) == 3


def test_manual_emotion_column_wins_and_is_checked(tmp_path: Path, tags_xlsx: Path) -> None:
    data = _csv(tmp_path / "d.csv", [{"text": "হ্যাঁ।", "emotion": "warm"}])
    _, rows = build_rows(data, load_provider_tags(tags_xlsx))
    assert rows[0][:2] == ["warm", "[warmly] হ্যাঁ।"]

    bad = _csv(tmp_path / "bad.csv", [{"text": "হ্যাঁ।", "emotion": "joyful"}])
    with pytest.raises(StyleMapError, match="not a Common Tag"):
        build_rows(bad, load_provider_tags(tags_xlsx))


def test_cli_writes_excel_that_opens(tmp_path: Path, tags_xlsx: Path) -> None:
    data = _csv(tmp_path / "d.csv", [{"text": "ধন্যবাদ। <&>"}, {"text": "হ্যাঁ।"}])
    out = tmp_path / "out.xlsx"
    assert main(["--csv", str(data), "--tags", str(tags_xlsx), "--out", str(out)]) == 0
    ws = load_workbook(out).active
    values = [[c.value for c in row] for row in ws.iter_rows()]
    assert values[0] == ["common tag", "Eleven Labs", "Gemini 3.8 Flash", "Cartesia Sonic 3", "Soniox"]
    assert values[1][:2] == ["warm", "[warmly] ধন্যবাদ। <&>"]
    assert values[2] == [NEUTRAL, "হ্যাঁ।", "হ্যাঁ।", "হ্যাঁ।", "হ্যাঁ।"]
