from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

REQUIRED_COLUMNS = ("id", "text", "spoken_form", "category")
SHEET = "Dataset"


class DatasetError(ValueError):
    pass


@dataclass(frozen=True)
class DatasetItem:
    item_id: str
    text: str  # sent to TTS byte-for-byte; never strip or normalize
    spoken_form: str | None
    accepted_variants: tuple[str, ...]
    category: str
    subcategory: str | None
    domain: str | None
    difficulty: int | None


def load_dataset(path: str | Path) -> list[DatasetItem]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise DatasetError(f"{path}: sheet {SHEET!r} not found (have {wb.sheetnames})")
    rows = wb[SHEET].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        raise DatasetError(f"{path}: missing required columns {missing}")
    col = {name: i for i, name in enumerate(header)}

    def cell(row: tuple, name: str):
        i = col.get(name)
        return row[i] if i is not None and i < len(row) else None

    items: list[DatasetItem] = []
    seen: set[str] = set()
    for line_no, row in enumerate(rows, start=2):
        item_id = cell(row, "id")
        if item_id is None or str(item_id).strip() == "":
            continue
        item_id = str(item_id).strip()
        text = cell(row, "text")
        if not isinstance(text, str) or text.strip() == "":
            raise DatasetError(f"{path}: row {line_no} ({item_id}) has empty or non-text 'text'")
        if item_id in seen:
            raise DatasetError(f"{path}: duplicate id {item_id} at row {line_no}")
        seen.add(item_id)
        variants = cell(row, "accepted_variants")
        difficulty = cell(row, "difficulty")
        items.append(
            DatasetItem(
                item_id=item_id,
                text=text,
                spoken_form=cell(row, "spoken_form"),
                accepted_variants=tuple(v for v in str(variants).split("|") if v) if variants else (),
                category=str(cell(row, "category")),
                subcategory=cell(row, "subcategory"),
                domain=cell(row, "domain"),
                difficulty=int(difficulty) if isinstance(difficulty, int | float) else None,
            )
        )
    wb.close()
    if not items:
        raise DatasetError(f"{path}: no rows in {SHEET}")
    return sorted(items, key=lambda it: it.item_id)
